#!/usr/bin/env python3
"""Gera os lançamentos financeiros da Trattoria Bela Serra que a página oferece para download.

Fonte da verdade: a constante D do index.html. O script monta os lançamentos que formam cada
número do DRE (pela data de competência) e do fluxo de caixa (pela data de pagamento), confere
tudo centavo a centavo e grava:

    dados/trattoria-lancamentos.xlsx   Leia-me, Lançamentos e as abas DRE e Fluxo de caixa,
                                       montadas com fórmulas (SOMASES) sobre os lançamentos
    dados/trattoria-lancamentos.csv    só os lançamentos, no padrão do Excel em português
                                       (ponto e vírgula, vírgula decimal, UTF-8)

Uso, na raiz do repositório:

    python3 ferramentas/gerar_lancamentos.py

Requer openpyxl (pip install openpyxl). Rode de novo sempre que mudar os números da página.
As sementes são fixas: sem mudança nos números, os lançamentos saem iguais. Se algum número
deixar de seguir as regras do caso, o script para e diz qual.

Regras do caso (as mesmas da nota da aba Fluxo de caixa):
- Salão: 40% das vendas entram no mesmo dia (dinheiro, PIX e débito). Os outros 60% são no
  crédito e caem no mesmo dia do mês seguinte, já sem a taxa da maquininha (3% das vendas).
- iFood: repassa no dia 10 do mês seguinte, já sem a comissão (23%) e o pagamento online (3,2%).
- Fornecedores (boleto para o mesmo dia do mês seguinte), Simples e folha: pagos no mês seguinte.
- Demais despesas, parcelas do empréstimo e retiradas: pagas no próprio mês.
Por isso o arquivo traz também lançamentos de ago/25 pagos em set/25 e lançamentos de ago/26
que só vencem em set/26 (em aberto na data dos dados).
"""

import calendar
import csv
import datetime as dt
import json
import math
import random
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parent.parent
PAGINA = RAIZ / 'index.html'
ARQ_XLSX = RAIZ / 'dados' / 'trattoria-lancamentos.xlsx'
ARQ_CSV = RAIZ / 'dados' / 'trattoria-lancamentos.csv'

SEMENTE = 'trattoria-bela-serra'
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']

# ---------- regras do caso ----------
A_VISTA = 0.40              # parte das vendas do salão que entra no mesmo dia
COMISSAO_IFOOD = 0.23       # + 3,2% de pagamento online = 26,2% do delivery
PARCELAS_EMPRESTIMO = 24
REFORMA = (dt.date(2025, 12, 1), dt.date(2025, 12, 10))   # salão fechado 10 dias
INICIO_DELIVERY = dt.date(2025, 12, 11)                    # reabertura, já com o iFood
FECHADO = {dt.date(2025, 12, 25), dt.date(2026, 1, 1)}     # além das segundas-feiras

# feriados nacionais sem expediente bancário (para as datas de pagamento)
FERIADOS = {dt.date(*d) for d in [
    (2025, 9, 7), (2025, 10, 12), (2025, 11, 2), (2025, 11, 15), (2025, 11, 20), (2025, 12, 25),
    (2026, 1, 1), (2026, 2, 16), (2026, 2, 17), (2026, 4, 3), (2026, 4, 21), (2026, 5, 1),
    (2026, 6, 4), (2026, 9, 7), (2026, 10, 12), (2026, 11, 2), (2026, 11, 15), (2026, 11, 20),
    (2026, 12, 25)]}

# peso de cada dia da semana nas vendas (ter a dom; segunda a casa fecha)
PESO_SALAO = {1: .75, 2: .80, 3: .90, 4: 1.25, 5: 1.55, 6: 1.50}
PESO_DELIVERY = {1: .85, 2: .90, 3: .95, 4: 1.20, 5: 1.30, 6: 1.25}
DATAS_FORTES = {dt.date(2025, 8, 10): 1.35, dt.date(2025, 12, 24): .8, dt.date(2025, 12, 31): .7,
                dt.date(2026, 5, 10): 1.6, dt.date(2026, 6, 12): 1.4, dt.date(2026, 8, 9): 1.35}

FORN = 'Fornecedores (insumos e embalagens)'
OUTROS_CX = 'Contabilidade, manutenção e outros'
OUTROS_DRE = 'Contabilidade, sistemas e outros'
FORA_DRE = 'Fora do DRE'

# plano de contas: categoria -> (grupo do DRE, linha do DRE, linha do fluxo de caixa, canal)
PLANO = {
    'Vendas salão (dinheiro, PIX e débito)': ('Receita bruta', 'Salão', 'Recebimentos do salão', 'Salão'),
    'Vendas salão (cartão de crédito)': ('Receita bruta', 'Salão', 'Recebimentos do salão', 'Salão'),
    'Vendas iFood': ('Receita bruta', 'Delivery (iFood)', 'Repasses do iFood', 'Delivery'),
    'Simples Nacional': ('Custos variáveis', 'Simples Nacional', 'Simples Nacional', 'Geral'),
    'Laticínios': ('Custos variáveis', 'CMV (insumos)', FORN, 'Geral'),
    'Carnes e aves': ('Custos variáveis', 'CMV (insumos)', FORN, 'Geral'),
    'Hortifrúti': ('Custos variáveis', 'CMV (insumos)', FORN, 'Geral'),
    'Mercearia (massas, grãos e secos)': ('Custos variáveis', 'CMV (insumos)', FORN, 'Geral'),
    'Bebidas': ('Custos variáveis', 'CMV (insumos)', FORN, 'Geral'),
    'Embalagens de delivery': ('Custos variáveis', 'Embalagens do delivery', FORN, 'Delivery'),
    'Comissão iFood (23%)': ('Custos variáveis', 'Comissão e taxas iFood', 'Repasses do iFood', 'Delivery'),
    'Pagamento online iFood (3,2%)': ('Custos variáveis', 'Comissão e taxas iFood', 'Repasses do iFood', 'Delivery'),
    'Taxas da maquininha (3%)': ('Custos variáveis', 'Taxas de cartão', 'Recebimentos do salão', 'Salão'),
    'Salários da cozinha': ('Despesas fixas', 'Folha e encargos', 'Folha e encargos', 'Geral'),
    'Salários do salão e administração': ('Despesas fixas', 'Folha e encargos', 'Folha e encargos', 'Geral'),
    'INSS e IRRF retidos': ('Despesas fixas', 'Folha e encargos', 'Folha e encargos', 'Geral'),
    'FGTS': ('Despesas fixas', 'Folha e encargos', 'Folha e encargos', 'Geral'),
    'Vale-transporte': ('Despesas fixas', 'Folha e encargos', 'Folha e encargos', 'Geral'),
    'Pró-labore': ('Despesas fixas', 'Pró-labore', 'Pró-labore', 'Geral'),
    'Aluguel': ('Despesas fixas', 'Aluguel', 'Aluguel', 'Geral'),
    'Energia elétrica': ('Despesas fixas', 'Energia, gás e água', 'Energia, gás e água', 'Geral'),
    'Gás': ('Despesas fixas', 'Energia, gás e água', 'Energia, gás e água', 'Geral'),
    'Água e esgoto': ('Despesas fixas', 'Energia, gás e água', 'Energia, gás e água', 'Geral'),
    'Agência de redes sociais': ('Despesas fixas', 'Marketing', 'Marketing', 'Geral'),
    'Anúncios em redes sociais': ('Despesas fixas', 'Marketing', 'Marketing', 'Geral'),
    'Contabilidade': ('Despesas fixas', OUTROS_DRE, OUTROS_CX, 'Geral'),
    'Sistema de PDV e gestão': ('Despesas fixas', OUTROS_DRE, OUTROS_CX, 'Geral'),
    'Internet e telefone': ('Despesas fixas', OUTROS_DRE, OUTROS_CX, 'Geral'),
    'Limpeza e higienização': ('Despesas fixas', OUTROS_DRE, OUTROS_CX, 'Geral'),
    'Lavanderia': ('Despesas fixas', OUTROS_DRE, OUTROS_CX, 'Geral'),
    'Controle de pragas': ('Despesas fixas', OUTROS_DRE, OUTROS_CX, 'Geral'),
    'Tarifas bancárias': ('Despesas fixas', OUTROS_DRE, OUTROS_CX, 'Geral'),
    'Manutenção e reparos': ('Despesas fixas', 'Manutenção', OUTROS_CX, 'Geral'),
    'Depreciação da reforma': ('Depreciação e juros', 'Depreciação da reforma', 'Fora do caixa', 'Geral'),
    'Juros do empréstimo': ('Depreciação e juros', 'Juros do empréstimo', 'Parcela do empréstimo', 'Geral'),
    'Empréstimo recebido': (FORA_DRE, FORA_DRE, 'Empréstimo recebido', 'Geral'),
    'Amortização do empréstimo': (FORA_DRE, FORA_DRE, 'Parcela do empréstimo', 'Geral'),
    'Reforma do salão': (FORA_DRE, FORA_DRE, 'Reforma do salão', 'Geral'),
    'Retirada dos sócios': (FORA_DRE, FORA_DRE, 'Retirada dos sócios', 'Geral'),
}
ORDEM_CATEGORIA = {c: i for i, c in enumerate(PLANO)}

# compras de insumos: categoria, fornecedor, nome na descrição, dias de entrega (0 = segunda),
# a cada quantas semanas, e o peso nas vendas do salão e do delivery
INSUMOS = [
    ('Laticínios', 'Laticínios Vale Serrano', 'laticínios', (1, 4), 1, .089, .120),
    ('Carnes e aves', 'Frigorífico Serra Nobre', 'carnes e aves', (2,), 1, .059, .062),
    ('Hortifrúti', 'Hortifrúti Bom Plantio', 'hortifrúti', (1, 3, 5), 1, .033, .034),
    ('Mercearia (massas, grãos e secos)', 'Empório Grão Fino', 'mercearia', (3,), 2, .050, .090),
    ('Bebidas', 'Adega & Cia Bebidas', 'bebidas', (2,), 2, .066, .020),
]
QUEIJO_NO_LATICINIO = .85   # quanto das compras de laticínios é muçarela e parmesão

# folha: estrutura antes e depois do 2º cozinheiro e do auxiliar (mar/26); soma = folha do mês
FOLHA_ANTES = [12480.00, 21420.00, 3330.00, 2978.40, 1791.60]     # 42.000
FOLHA_DEPOIS = [18530.00, 21420.00, 3890.00, 3507.20, 2652.80]    # 50.000

CONTAS_CONSUMO = [  # categoria, descrição, fornecedor, dia de vencimento, forma, peso
    ('Energia elétrica', 'Conta de energia elétrica', 'Distribuidora de energia', 15, 'Débito automático', .55),
    ('Gás', 'Gás de cozinha (GLP)', 'Gás Chama Azul', 20, 'Boleto', .30),
    ('Água e esgoto', 'Conta de água e esgoto', 'Companhia de saneamento', 12, 'Débito automático', .15),
]
OUTROS = [  # categoria, descrição, fornecedor, dia de vencimento, forma, valor do mês (soma 3.500)
    ('Contabilidade', 'Honorários contábeis', 'Escritório Contábil Exato', 10, 'Boleto', 1450.00),
    ('Sistema de PDV e gestão', 'Mensalidade do sistema de PDV e gestão', 'GestPDV Sistemas', 5, 'Débito automático', 389.90),
    ('Internet e telefone', 'Internet e telefone', 'NetSerra Telecom', 15, 'Débito automático', 249.90),
    ('Limpeza e higienização', 'Limpeza e higienização (contrato)', 'Limpa Bem Higienização', 5, 'Boleto', 780.00),
    ('Lavanderia', 'Lavanderia de toalhas e uniformes', 'Lavanderia Branco Neve', 25, 'PIX', 420.00),
    ('Controle de pragas', 'Controle de pragas (contrato)', 'Dedetizadora Sem Praga', 20, 'PIX', 150.00),
    ('Tarifas bancárias', 'Tarifas bancárias', 'Banco Horizonte', 1, 'Débito em conta', 60.20),
]
SERVICOS = [  # manutenção: descrição, fornecedor
    ('Manutenção preventiva da câmara fria', 'Polar Refrigeração'),
    ('Recarga de gás do ar-condicionado do salão', 'Polar Refrigeração'),
    ('Conserto do forno combinado', 'TecForno Assistência Técnica'),
    ('Manutenção do cilindro de massas', 'TecForno Assistência Técnica'),
    ('Manutenção da máquina de lavar louça', 'TecForno Assistência Técnica'),
    ('Limpeza da coifa e do exaustor', 'Coifa Limpa Serviços'),
    ('Reparo hidráulico na cozinha', 'Hidráulica Serra'),
    ('Desentupimento da caixa de gordura', 'Hidráulica Serra'),
    ('Reparo elétrico no salão', 'Eletricista Souza'),
    ('Afiação de facas e reparo de utensílios', 'Afiação Fio Certo'),
    ('Conserto de mesas e cadeiras', 'Marcenaria Pinho Nobre'),
]
OBRA = [  # reforma do salão: descrição, fornecedor, dia do pagamento, peso
    ('Reforma do salão: obra civil e acabamento', 'Construtora Alicerce', 1, 82000),
    ('Reforma do salão: marcenaria e balcão', 'Marcenaria Pinho Nobre', 3, 34500),
    ('Reforma do salão: elétrica e iluminação', 'Luz & Forma Instalações', 5, 16800),
    ('Reforma do salão: mesas e cadeiras', 'Casa Rústica Móveis', 10, 16700),
]
SOCIOS = ['Marco', 'Júlia']


# ---------- utilidades ----------
def centavos(v):
    return round(v * 100)


def alocar(total, pesos):
    """Divide `total` (em centavos) na proporção dos pesos; a soma das partes é exata."""
    soma = sum(pesos)
    brutos = [total * p / soma for p in pesos]
    partes = [math.floor(b) for b in brutos]
    ordem = sorted(range(len(pesos)), key=lambda i: brutos[i] - partes[i], reverse=True)
    for i in ordem[:total - sum(partes)]:
        partes[i] += 1
    return partes


def sorteio(*chave):
    return random.Random('-'.join(map(str, (SEMENTE,) + chave)))


def rotulo(ano, mes):
    return f'{MESES[mes - 1]}/{ano % 100:02d}'


def seguinte(ano, mes):
    return (ano + 1, 1) if mes == 12 else (ano, mes + 1)


def anterior(ano, mes):
    return (ano - 1, 12) if mes == 1 else (ano, mes - 1)


def dias_do_mes(ano, mes):
    return [dt.date(ano, mes, d) for d in range(1, calendar.monthrange(ano, mes)[1] + 1)]


def ultimo_dia(ano, mes):
    return dt.date(ano, mes, calendar.monthrange(ano, mes)[1])


def util(d):
    return d.weekday() < 5 and d not in FERIADOS


def ajusta_util(d):
    """Dia útil seguinte; se ele cair no outro mês, o dia útil anterior."""
    x = d
    while not util(x):
        x += dt.timedelta(days=1)
    return x if x.month == d.month else util_anterior(d)


def util_anterior(d):
    while not util(d):
        d -= dt.timedelta(days=1)
    return d


def dia_util(ano, mes, n):
    """n-ésimo dia útil do mês (n = -1: o último)."""
    uteis = [d for d in dias_do_mes(ano, mes) if util(d)]
    return uteis[n - 1] if n > 0 else uteis[n]


def mais_um_mes(d):
    ano, mes = seguinte(d.year, d.month)
    return dt.date(ano, mes, min(d.day, calendar.monthrange(ano, mes)[1]))


def aberto(d):
    return d.weekday() != 0 and d not in FECHADO and not REFORMA[0] <= d <= REFORMA[1]


def ddmmaa(d):
    return d.strftime('%d/%m/%y')


# ---------- lançamentos ----------
LANCAMENTOS = []
FIM = None   # último dia dos dados (31/08/26): o que vence depois fica em aberto


def lanca(comp, venc, pag, descricao, contraparte, categoria, valor, forma):
    grupo, linha_dre, linha_fluxo, canal = PLANO[categoria]
    if pag and pag > FIM:
        pag = None
    LANCAMENTOS.append(dict(comp=comp, venc=venc, pag=pag, descricao=descricao, contraparte=contraparte,
                            categoria=categoria, canal=canal, grupo=grupo, linha_dre=linha_dre,
                            linha_fluxo=linha_fluxo, forma=forma, valor=valor))


def vendas_salao(ano, mes, total, taxa):
    r = sorteio('salao', ano, mes)
    dias = [d for d in dias_do_mes(ano, mes) if aberto(d)]
    vendas = alocar(centavos(total), [PESO_SALAO[d.weekday()] * DATAS_FORTES.get(d, 1) * r.uniform(.88, 1.12) for d in dias])
    avista = alocar(round(centavos(total) * A_VISTA), [v * r.uniform(.34, .46) for v in vendas])
    taxas = alocar(centavos(taxa), vendas)
    for d, v, a, t in zip(dias, vendas, avista, taxas):
        lanca(d, d, d, f'Vendas do salão em {ddmmaa(d)}: dinheiro, PIX e débito', 'Clientes do salão',
              'Vendas salão (dinheiro, PIX e débito)', a, 'Dinheiro, PIX e débito')
        venc = mais_um_mes(d)
        lanca(d, venc, ajusta_util(venc), f'Vendas do salão em {ddmmaa(d)}: cartão de crédito', 'Operadora da maquininha',
              'Vendas salão (cartão de crédito)', v - a, 'Cartão de crédito')
        lanca(d, venc, ajusta_util(venc), f'Taxa da maquininha sobre as vendas de {ddmmaa(d)}', 'Operadora da maquininha',
              'Taxas da maquininha (3%)', -t, 'Descontada no repasse do cartão')


def vendas_ifood(ano, mes, total, taxas):
    if not total:
        return
    r = sorteio('ifood', ano, mes)
    dias = [d for d in dias_do_mes(ano, mes) if aberto(d) and d >= INICIO_DELIVERY]
    # nas primeiras semanas o delivery ainda estava ganhando pedidos
    pesos = [PESO_DELIVERY[d.weekday()] * min(1, .45 + (d - INICIO_DELIVERY).days / 36) * r.uniform(.85, 1.15) for d in dias]
    venc = dt.date(*seguinte(ano, mes), 10)
    pag = ajusta_util(venc)
    for d, v in zip(dias, alocar(centavos(total), pesos)):
        lanca(d, venc, pag, f'Pedidos do iFood em {ddmmaa(d)}', 'iFood', 'Vendas iFood', v, 'Repasse do iFood')
    comissao = round(centavos(total) * COMISSAO_IFOOD)
    fim, rot = ultimo_dia(ano, mes), rotulo(ano, mes)
    lanca(fim, venc, pag, f'Comissão do iFood (23%) sobre os pedidos de {rot}', 'iFood', 'Comissão iFood (23%)',
          -comissao, 'Descontada no repasse do iFood')
    lanca(fim, venc, pag, f'Pagamento online do iFood (3,2%) sobre os pedidos de {rot}', 'iFood',
          'Pagamento online iFood (3,2%)', -(centavos(taxas) - comissao), 'Descontada no repasse do iFood')


NOTA_FISCAL = {}


def notas(r, ano, mes, categoria, fornecedor, nome, dias_semana, cada, total):
    """Compras do mês em notas fiscais nos dias de entrega, com boleto para o mesmo dia do mês seguinte."""
    datas = [d for d in dias_do_mes(ano, mes)
             if d.weekday() in dias_semana and d not in FERIADOS and not REFORMA[0] <= d < REFORMA[1]][::cada]
    if fornecedor not in NOTA_FISCAL:
        NOTA_FISCAL[fornecedor] = sorteio('nf', fornecedor).randint(2000, 60000)
    for d, v in zip(datas, alocar(total, [r.uniform(.75, 1.25) for _ in datas])):
        NOTA_FISCAL[fornecedor] += r.randint(6, 90)
        venc = mais_um_mes(d)
        lanca(d, venc, ajusta_util(venc), f'Compra de {nome}: NF {NOTA_FISCAL[fornecedor]}', fornecedor, categoria, -v, 'Boleto')


def compras(ano, mes, v, preco_queijo):
    r = sorteio('compras', ano, mes)
    indice = QUEIJO_NO_LATICINIO * preco_queijo + (1 - QUEIJO_NO_LATICINIO)
    pesos = [(ps * v['salao'] + pd * v['delivery']) * (indice if cat == 'Laticínios' else 1)
             for cat, _, _, _, _, ps, pd in INSUMOS]
    for (cat, forn, nome, dias, cada, _, _), total in zip(INSUMOS, alocar(centavos(v['cmv']), pesos)):
        notas(r, ano, mes, cat, forn, nome, dias, cada, total)
    if v['embalagem']:
        notas(r, ano, mes, 'Embalagens de delivery', 'Pack Mais Embalagens', 'embalagens de delivery', (2,), 2,
              centavos(v['embalagem']))


def simples(ano, mes, valor):
    venc = dt.date(*seguinte(ano, mes), 20)
    lanca(ultimo_dia(ano, mes), venc, ajusta_util(venc), f'DAS do Simples Nacional, competência {rotulo(ano, mes)}',
          'Receita Federal', 'Simples Nacional', -centavos(valor), 'Guia DAS')


def folha(ano, mes, valor, folha_inicial):
    fim, rot, (a2, m2) = ultimo_dia(ano, mes), rotulo(ano, mes), seguinte(ano, mes)
    quinto, dia20 = dia_util(a2, m2, 5), dt.date(a2, m2, 20)
    partes = alocar(centavos(valor), FOLHA_DEPOIS if valor > folha_inicial else FOLHA_ANTES)
    itens = [
        ('Salários da cozinha', f'Salários da cozinha, folha de {rot}', 'Funcionários da cozinha', quinto, quinto, 'Transferência bancária'),
        ('Salários do salão e administração', f'Salários do salão e da administração, folha de {rot}',
         'Funcionários do salão e da administração', quinto, quinto, 'Transferência bancária'),
        ('INSS e IRRF retidos', f'INSS e IRRF retidos dos funcionários, folha de {rot}', 'Receita Federal',
         dia20, util_anterior(dia20), 'Guia DCTFWeb'),
        ('FGTS', f'FGTS, folha de {rot}', 'FGTS (Caixa)', dia20, util_anterior(dia20), 'Guia FGTS Digital'),
        ('Vale-transporte', f'Vale-transporte, folha de {rot}', 'Operadora de vale-transporte', quinto, quinto, 'Boleto'),
    ]
    for (cat, desc, quem, venc, pag, forma), v in zip(itens, partes):
        lanca(fim, venc, pag, desc, quem, cat, -v, forma)


def pro_labore(ano, mes, valor):
    d = dia_util(ano, mes, -1)
    for socio, v in zip(SOCIOS, alocar(centavos(valor), [1] * len(SOCIOS))):
        lanca(d, d, d, f'Pró-labore de {socio}, {rotulo(ano, mes)}', socio, 'Pró-labore', -v, 'PIX')


def aluguel(ano, mes, valor):
    venc = dt.date(ano, mes, 10)
    lanca(venc, venc, ajusta_util(venc), f'Aluguel do imóvel, {rotulo(ano, mes)}', 'Imobiliária Pedra Alta',
          'Aluguel', -centavos(valor), 'Boleto')


def contas_do_mes(ano, mes, itens, valores):
    for (cat, desc, quem, dia, forma, _), v in zip(itens, valores):
        venc = dt.date(ano, mes, dia)
        lanca(venc, venc, ajusta_util(venc), f'{desc}, {rotulo(ano, mes)}', quem, cat, -v, forma)


def energia(ano, mes, valor):
    r = sorteio('energia', ano, mes)
    contas_do_mes(ano, mes, CONTAS_CONSUMO, alocar(centavos(valor), [p * r.uniform(.93, 1.07) for *_, p in CONTAS_CONSUMO]))


def marketing(ano, mes, valor, base):
    agencia = min(valor, base)
    venc = dt.date(ano, mes, 5)
    lanca(venc, venc, ajusta_util(venc), f'Gestão das redes sociais, {rotulo(ano, mes)}', 'Agência Manjericão Digital',
          'Agência de redes sociais', -centavos(agencia), 'Boleto')
    if valor > agencia:
        d = dia_util(ano, mes, 1)
        lanca(d, d, d, f'Anúncios no Instagram e no Facebook, {rotulo(ano, mes)}', 'Meta (anúncios)',
              'Anúncios em redes sociais', -centavos(valor - agencia), 'PIX')


def outros(ano, mes, valor):
    contas_do_mes(ano, mes, OUTROS, alocar(centavos(valor), [v for *_, v in OUTROS]))


def manutencao(ano, mes, valor):
    r = sorteio('manutencao', ano, mes)
    n = r.choice([2, 2, 3])
    datas = sorted(r.sample([d for d in dias_do_mes(ano, mes) if util(d)], n))
    for (desc, quem), d, v in zip(r.sample(SERVICOS, n), datas, alocar(centavos(valor), [r.uniform(.5, 1.5) for _ in range(n)])):
        lanca(d, d, d, desc, quem, 'Manutenção e reparos', -v, 'PIX')


def depreciacao(ano, mes, valor, texto):
    lanca(ultimo_dia(ano, mes), None, None, f'Depreciação da reforma do salão{texto}, {rotulo(ano, mes)}', '',
          'Depreciação da reforma', -centavos(valor), '')


def emprestimo(ano, mes, valor):
    d = ajusta_util(dt.date(ano, mes, 14))
    lanca(d, d, d, f'Empréstimo para a reforma do salão ({PARCELAS_EMPRESTIMO} parcelas)', 'Banco Horizonte',
          'Empréstimo recebido', centavos(valor), 'Crédito em conta')


def parcela(ano, mes, valor, juros, n):
    venc = dt.date(ano, mes, 15)
    pag = ajusta_util(venc)
    txt = f'Parcela {n}/{PARCELAS_EMPRESTIMO} do empréstimo'
    lanca(venc, venc, pag, f'{txt}: juros', 'Banco Horizonte', 'Juros do empréstimo', -centavos(juros), 'Débito em conta')
    lanca(venc, venc, pag, f'{txt}: amortização', 'Banco Horizonte', 'Amortização do empréstimo',
          -centavos(valor - juros), 'Débito em conta')


def reforma(ano, mes, valor):
    for (desc, quem, dia, _), v in zip(OBRA, alocar(centavos(valor), [p for *_, p in OBRA])):
        d = ajusta_util(dt.date(ano, mes, dia))
        lanca(d, d, d, desc, quem, 'Reforma do salão', -v, 'Transferência bancária')


def retirada(ano, mes, valor):
    d = dia_util(ano, mes, 5)
    for socio, v in zip(SOCIOS, alocar(centavos(valor), [1] * len(SOCIOS))):
        lanca(d, d, d, f'Retirada de {socio}, {rotulo(ano, mes)}', socio, 'Retirada dos sócios', -v, 'PIX')


# ---------- leitura da página e montagem ----------
def le_pagina():
    m = re.search(r'^const D = (\{.*\});$', PAGINA.read_text(encoding='utf-8'), re.M)
    if not m:
        sys.exit('Não achei a linha "const D = {...};" em index.html')
    return json.loads(m.group(1))


def mes_da_pagina(txt):
    nome, aa = txt.split('/')
    return 2000 + int(aa), MESES.index(nome) + 1


def falha(msg):
    sys.exit(f'ERRO: {msg}')


def monta(D):
    global FIM
    meses = [mes_da_pagina(d['mes']) for d in D['dre']]
    FIM = ultimo_dia(*meses[-1])
    ini = dt.date(*meses[0], 1)

    # mês anterior ao caso: sai das regras aplicadas ao caixa do primeiro mês
    c0, d0 = D['caixa'][0], D['dre'][0]
    salao_ant = (c0['ent_salao'] - A_VISTA * d0['salao']) / (1 - A_VISTA - d0['cartao'] / d0['salao'])
    if abs(salao_ant - round(salao_ant)) > 1e-6 or c0['ent_ifood']:
        falha('o caixa do primeiro mês não segue as regras de recebimento do salão e do iFood')
    ant = dict(salao=round(salao_ant), delivery=0, cmv=c0['saidas']['fornecedores'], embalagem=0,
               cartao=round(round(salao_ant) * d0['cartao'] / d0['salao']))
    a0 = anterior(*meses[0])

    # preço dos queijos: estável até o primeiro mês do gráfico
    q = D['queijo']
    fator = {m: (q['mucarela'][i] / q['mucarela'][0] + q['parmesao'][i] / q['parmesao'][0]) / 2
             for i, m in enumerate(q['meses'])}
    queijo = [fator[MESES[m - 1]] if i >= len(meses) - len(q['meses']) else 1 for i, (_, m) in enumerate(meses)]

    reforma_total = sum(c['saidas']['reforma'] for c in D['caixa'])
    deprec = next((d['deprec'] for d in D['dre'] if d['deprec']), 0)
    anos = reforma_total / deprec / 12 if deprec else 0
    texto_deprec = f' (R$ {reforma_total // 1000} mil em {anos:.0f} anos)' if anos == int(anos) and anos else ''

    # mês anterior: só entra o que foi pago dentro do período do caso
    vendas_salao(*a0, ant['salao'], ant['cartao'])
    compras(*a0, ant, 1)
    simples(*a0, c0['saidas']['simples'])
    folha(*a0, c0['saidas']['folha'], c0['saidas']['folha'])

    n_parcela = 0
    for (ano, mes), d, c, pq in zip(meses, D['dre'], D['caixa'], queijo):
        vendas_salao(ano, mes, d['salao'], d['cartao'])
        vendas_ifood(ano, mes, d['delivery'], d['ifood'])
        compras(ano, mes, d, pq)
        simples(ano, mes, d['simples'])
        folha(ano, mes, d['folha'], D['dre'][0]['folha'])
        pro_labore(ano, mes, d['prolabore'])
        aluguel(ano, mes, d['aluguel'])
        energia(ano, mes, d['energia'])
        marketing(ano, mes, d['marketing'], D['dre'][0]['marketing'])
        outros(ano, mes, d['outros'])
        manutencao(ano, mes, d['manutencao'])
        if d['deprec']:
            depreciacao(ano, mes, d['deprec'], texto_deprec)
        if c['emprestimo']:
            emprestimo(ano, mes, c['emprestimo'])
        if c['saidas']['parcela']:
            n_parcela += 1
            parcela(ano, mes, c['saidas']['parcela'], d['juros'], n_parcela)
        if c['saidas']['reforma']:
            reforma(ano, mes, c['saidas']['reforma'])
        if c['saidas']['retirada']:
            retirada(ano, mes, c['saidas']['retirada'])

    dentro = lambda x: x is not None and ini <= x <= FIM
    lista = [l for l in LANCAMENTOS if dentro(l['comp']) or dentro(l['pag'])]
    lista.sort(key=lambda l: (l['comp'], l['valor'] < 0, ORDEM_CATEGORIA[l['categoria']], l['descricao']))
    for i, l in enumerate(lista, 1):
        l['n'] = i
        l['mes_comp'] = l['comp'].replace(day=1)
        l['mes_caixa'] = l['pag'].replace(day=1) if l['pag'] else None
        if l['linha_fluxo'] == 'Fora do caixa':
            l['situacao'] = 'Sem efeito no caixa'
        elif l['valor'] > 0:
            l['situacao'] = 'Recebido' if l['pag'] else 'A receber'
        else:
            l['situacao'] = 'Pago' if l['pag'] else 'A pagar'
    return meses, lista


# ---------- conferência com a página ----------
DRE_PAGINA = [  # linha do DRE nos lançamentos, chave no D, sinal
    ('Salão', 'salao', 1), ('Delivery (iFood)', 'delivery', 1), ('Simples Nacional', 'simples', -1),
    ('CMV (insumos)', 'cmv', -1), ('Embalagens do delivery', 'embalagem', -1), ('Comissão e taxas iFood', 'ifood', -1),
    ('Taxas de cartão', 'cartao', -1), ('Folha e encargos', 'folha', -1), ('Pró-labore', 'prolabore', -1),
    ('Aluguel', 'aluguel', -1), ('Energia, gás e água', 'energia', -1), ('Marketing', 'marketing', -1),
    (OUTROS_DRE, 'outros', -1), ('Manutenção', 'manutencao', -1), ('Depreciação da reforma', 'deprec', -1),
    ('Juros do empréstimo', 'juros', -1),
]
FLUXO_PAGINA = [  # linha do fluxo nos lançamentos, chave no D (saídas ficam em D.caixa[i].saidas), sinal
    ('Recebimentos do salão', 'ent_salao', 1), ('Repasses do iFood', 'ent_ifood', 1), ('Empréstimo recebido', 'emprestimo', 1),
    (FORN, 'fornecedores', -1), ('Simples Nacional', 'simples', -1), ('Folha e encargos', 'folha', -1),
    ('Pró-labore', 'prolabore', -1), ('Aluguel', 'aluguel', -1), ('Energia, gás e água', 'energia', -1),
    ('Marketing', 'marketing', -1), (OUTROS_CX, 'outros', -1), ('Reforma do salão', 'reforma', -1),
    ('Parcela do empréstimo', 'parcela', -1), ('Retirada dos sócios', 'retirada', -1),
]


def soma(lista, campo_mes, campo_linha, mes, linha):
    return sum(l['valor'] for l in lista if l[campo_mes] == mes and l[campo_linha] == linha)


def confere(D, meses, lista):
    erros = []
    for (ano, mes), d, c in zip(meses, D['dre'], D['caixa']):
        m = dt.date(ano, mes, 1)
        for linha, chave, sinal in DRE_PAGINA:
            if soma(lista, 'mes_comp', 'linha_dre', m, linha) != sinal * centavos(d[chave]):
                erros.append(f'DRE {d["mes"]} · {linha}')
        if sum(l['valor'] for l in lista if l['mes_comp'] == m and l['linha_dre'] != FORA_DRE) != centavos(d['lucro']):
            erros.append(f'DRE {d["mes"]} · lucro líquido')
        for linha, chave, sinal in FLUXO_PAGINA:
            alvo = c[chave] if chave in c else c['saidas'][chave]
            if soma(lista, 'mes_caixa', 'linha_fluxo', m, linha) != sinal * centavos(alvo):
                erros.append(f'Fluxo {d["mes"]} · {linha}')
        if sum(l['valor'] for l in lista if l['mes_caixa'] == m) != centavos(c['resultado']):
            erros.append(f'Fluxo {d["mes"]} · resultado do mês')
    if erros:
        falha('os lançamentos não batem com a página em:\n  ' + '\n  '.join(erros))


# ---------- arquivos ----------
COLUNAS = [  # título, campo, largura, formato
    ('Nº', 'n', 7, '0'),
    ('Data de competência', 'comp', 13, 'dd/mm/yyyy'),
    ('Data de vencimento', 'venc', 13, 'dd/mm/yyyy'),
    ('Data de pagamento', 'pag', 13, 'dd/mm/yyyy'),
    ('Situação', 'situacao', 12, None),
    ('Descrição', 'descricao', 54, None),
    ('Cliente/fornecedor', 'contraparte', 30, None),
    ('Categoria', 'categoria', 33, None),
    ('Canal', 'canal', 10, None),
    ('Grupo do DRE', 'grupo', 19, None),
    ('Linha do DRE', 'linha_dre', 31, None),
    ('Linha do fluxo de caixa', 'linha_fluxo', 34, None),
    ('Forma de pagamento', 'forma', 31, None),
    ('Valor (R$)', 'valor', 14, '#,##0.00;[Red]-#,##0.00'),
    ('Mês de competência', 'mes_comp', 12, 'mmm/yy'),
    ('Mês de caixa', 'mes_caixa', 12, 'mmm/yy'),
]
COL = {campo: chr(ord('A') + i) for i, (_, campo, _, _) in enumerate(COLUNAS)}

EXPLICA_COLUNAS = {
    'Nº': 'Número do lançamento.',
    'Data de competência': 'Quando a venda ou a despesa aconteceu. É a data que conta no DRE.',
    'Data de vencimento': 'Quando o valor vence ou venceu.',
    'Data de pagamento': 'Quando o dinheiro entrou ou saiu do banco. É a data que conta no fluxo de caixa. '
                         'Fica vazia no que ainda está em aberto e na depreciação.',
    'Situação': 'Recebido, Pago, A receber, A pagar ou Sem efeito no caixa (depreciação).',
    'Descrição': 'O que é o lançamento.',
    'Cliente/fornecedor': 'De quem vem ou para quem vai o dinheiro.',
    'Categoria': 'Classificação detalhada do lançamento.',
    'Canal': 'Salão ou Delivery quando o lançamento é de um canal só; Geral quando é da casa toda.',
    'Grupo do DRE': 'Receita bruta, Custos variáveis, Despesas fixas, Depreciação e juros ou Fora do DRE.',
    'Linha do DRE': 'Linha do DRE da página onde o lançamento entra. Fora do DRE: não é receita nem despesa '
                    '(empréstimo, amortização, reforma e retiradas).',
    'Linha do fluxo de caixa': 'Linha do fluxo de caixa da página onde o lançamento entra. '
                               'Fora do caixa: não passa pelo banco (depreciação).',
    'Forma de pagamento': 'Como o valor foi ou será pago ou recebido.',
    'Valor (R$)': 'Positivo para receitas e entradas de dinheiro; negativo para despesas e saídas.',
    'Mês de competência': 'Mês da data de competência, para agrupar em tabelas dinâmicas.',
    'Mês de caixa': 'Mês da data de pagamento, para agrupar em tabelas dinâmicas.',
}

# layout das abas de resumo: rótulo, estilo (g grupo, f linha, r resultado, t total), linha nos lançamentos ou soma
DRE_ABA = [
    ('Receita bruta', 'g', ['Salão', 'Delivery (iFood)']),
    ('Salão', 'f', 'Salão'), ('Delivery (iFood)', 'f', 'Delivery (iFood)'),
    ('(-) Custos variáveis', 'g', ['Simples Nacional', 'CMV (insumos)', 'Embalagens do delivery',
                                   'Comissão e taxas iFood', 'Taxas de cartão']),
    ('Simples Nacional', 'f', 'Simples Nacional'), ('CMV (insumos)', 'f', 'CMV (insumos)'),
    ('Embalagens do delivery', 'f', 'Embalagens do delivery'), ('Comissão e taxas iFood', 'f', 'Comissão e taxas iFood'),
    ('Taxas de cartão', 'f', 'Taxas de cartão'),
    ('Margem de contribuição', 'r', ['Receita bruta', '(-) Custos variáveis']),
    ('(-) Despesas fixas', 'g', ['Folha e encargos', 'Pró-labore', 'Aluguel', 'Energia, gás e água', 'Marketing',
                                 OUTROS_DRE, 'Manutenção']),
    ('Folha e encargos', 'f', 'Folha e encargos'), ('Pró-labore', 'f', 'Pró-labore'), ('Aluguel', 'f', 'Aluguel'),
    ('Energia, gás e água', 'f', 'Energia, gás e água'), ('Marketing', 'f', 'Marketing'),
    (OUTROS_DRE, 'f', OUTROS_DRE), ('Manutenção', 'f', 'Manutenção'),
    ('EBITDA', 'r', ['Margem de contribuição', '(-) Despesas fixas']),
    ('(-) Depreciação da reforma', 'f', 'Depreciação da reforma'), ('(-) Juros do empréstimo', 'f', 'Juros do empréstimo'),
    ('Lucro líquido', 't', ['EBITDA', '(-) Depreciação da reforma', '(-) Juros do empréstimo']),
]
FLUXO_ABA = [
    ('Saldo inicial', 'r', None),
    ('Entradas', 'g', ['Recebimentos do salão', 'Repasses do iFood', 'Empréstimo recebido']),
    ('Recebimentos do salão', 'f', 'Recebimentos do salão'), ('Repasses do iFood', 'f', 'Repasses do iFood'),
    ('Empréstimo recebido', 'f', 'Empréstimo recebido'),
    ('Saídas', 'g', [FORN, 'Simples Nacional', 'Folha e encargos', 'Pró-labore', 'Aluguel', 'Energia, gás e água',
                     'Marketing', OUTROS_CX, 'Reforma do salão', 'Parcela do empréstimo', 'Retirada dos sócios']),
    (FORN, 'f', FORN), ('Simples Nacional', 'f', 'Simples Nacional'), ('Folha e encargos', 'f', 'Folha e encargos'),
    ('Pró-labore', 'f', 'Pró-labore'), ('Aluguel', 'f', 'Aluguel'), ('Energia, gás e água', 'f', 'Energia, gás e água'),
    ('Marketing', 'f', 'Marketing'), (OUTROS_CX, 'f', OUTROS_CX), ('Reforma do salão', 'f', 'Reforma do salão'),
    ('Parcela do empréstimo', 'f', 'Parcela do empréstimo'), ('Retirada dos sócios', 'f', 'Retirada dos sócios'),
    ('Resultado do mês', 'r', ['Entradas', 'Saídas']),
    ('Saldo final', 't', None),
]

TEAL, TEAL_LEVE, FUNDO, SUAVE = '0C4448', 'E7F1EF', 'F7F9F9', '6A787A'


def fonte(**kw):
    return Font(name='Arial', size=kw.pop('size', 10), **kw)


def preenche(cor):
    return PatternFill('solid', start_color=cor, end_color=cor)


def aba_lancamentos(wb, lista):
    ws = wb.create_sheet('Lançamentos')
    for j, (titulo, _, largura, _) in enumerate(COLUNAS, 1):
        c = ws.cell(1, j, titulo)
        c.font = fonte(bold=True, color='FFFFFF')
        c.fill = preenche(TEAL)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[chr(64 + j)].width = largura
    ws.row_dimensions[1].height = 30
    for i, l in enumerate(lista, 2):
        for j, (_, campo, _, formato) in enumerate(COLUNAS, 1):
            v = l[campo]
            c = ws.cell(i, j, v / 100 if campo == 'valor' else v)
            c.font = fonte()
            if formato:
                c.number_format = formato
    tab = Table(displayName='Lancamentos', ref=f'A1:{chr(64 + len(COLUNAS))}{len(lista) + 1}')
    tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = 'A2'


def aba_resumo(wb, nome, titulo, explicacao, layout, meses, lista, campo_mes, campo_linha, saldo_inicial=None):
    """Aba com o DRE ou o fluxo de caixa refeitos por fórmulas (SOMASES) sobre a aba Lançamentos.
    Devolve {célula: resultado esperado, em centavos} para gravar junto com as fórmulas."""
    ws = wb.create_sheet(nome)
    ultima = len(lista) + 1
    faixa = lambda campo: f"'Lançamentos'!${COL[campo]}$2:${COL[campo]}${ultima}"
    col = lambda j: chr(ord('B') + j)
    n = len(meses)
    ws['A1'] = titulo
    ws['A1'].font = fonte(bold=True, size=14, color=TEAL)
    ws['A2'] = explicacao
    ws['A2'].font = fonte(italic=True, color=SUAVE)
    ws.column_dimensions['A'].width = 38
    ws.cell(4, 1, 'R$')
    for j, (ano, mes) in enumerate(meses):
        ws.cell(4, 2 + j, dt.date(ano, mes, 1)).number_format = 'mmm/yy'
        ws.column_dimensions[col(j)].width = 11
    ws.cell(4, 2 + n, 'Total')
    ws.column_dimensions[col(n)].width = 12
    for c in ws[4]:
        c.font = fonte(bold=True, color='FFFFFF')
        c.fill = preenche(TEAL)
        c.alignment = Alignment(horizontal='left' if c.column == 1 else 'right')

    linha_de = {rot: 5 + i for i, (rot, _, _) in enumerate(layout)}
    ref_de = {rot: ref for rot, _, ref in layout}
    valores = {}

    def valor(rot, j):   # em centavos; um grupo vem antes das suas linhas, por isso a recursão
        if (rot, j) not in valores:
            ref = ref_de[rot]
            if isinstance(ref, str):
                v = soma(lista, campo_mes, campo_linha, dt.date(*meses[j], 1), ref)
            elif ref:
                v = sum(valor(x, j) for x in ref)
            elif rot == 'Saldo inicial':
                v = centavos(saldo_inicial) if j == 0 else valor('Saldo final', j - 1)
            else:
                v = valor('Saldo inicial', j) + valor('Resultado do mês', j)
            valores[(rot, j)] = v
        return valores[(rot, j)]

    esperado = {}
    for rot, estilo, ref in layout:
        r = linha_de[rot]
        ws.cell(r, 1, rot)
        for j in range(n):
            if isinstance(ref, str):
                f = f'=SUMIFS({faixa("valor")},{faixa(campo_linha)},"{ref}",{faixa(campo_mes)},{col(j)}$4)'
            elif ref:
                f = '=' + '+'.join(f'{col(j)}{linha_de[x]}' for x in ref)
            elif rot == 'Saldo inicial':
                f = saldo_inicial if j == 0 else f'={col(j - 1)}{linha_de["Saldo final"]}'
            else:
                f = f'={col(j)}{linha_de["Saldo inicial"]}+{col(j)}{linha_de["Resultado do mês"]}'
            ws[f'{col(j)}{r}'] = f
            if isinstance(f, str):
                esperado[f'{col(j)}{r}'] = valor(rot, j)
        if ref is not None:   # saldo não se soma no ano
            ws[f'{col(n)}{r}'] = f'=SUM(B{r}:{col(n - 1)}{r})'
            esperado[f'{col(n)}{r}'] = sum(valor(rot, j) for j in range(n))
        for c in ws[r]:
            c.font = fonte()
            if c.column > 1:
                c.number_format = '#,##0;[Red]-#,##0;"-"'
            if estilo == 'g':
                c.font, c.fill = fonte(bold=True, color=TEAL), preenche(TEAL_LEVE)
            elif estilo == 'f' and c.column == 1:
                c.font, c.alignment = fonte(color=SUAVE), Alignment(indent=1)
            elif estilo == 'r':
                c.font, c.fill = fonte(bold=True), preenche(FUNDO)
            elif estilo == 't':
                c.font, c.fill = fonte(bold=True, color='FFFFFF'), preenche(TEAL)
                if c.column > 1:
                    c.number_format = '#,##0;-#,##0;"-"'
    if saldo_inicial is not None:
        c = ws.cell(linha_de['Saldo inicial'], 2)
        c.font = fonte(bold=True, color='0000FF')
        c.comment = Comment(f'Saldo no banco em {dt.date(*meses[0], 1):%d/%m/%y}, informado na página. '
                            'Os outros meses partem do saldo final do mês anterior.', 'Zelo')
    ws.freeze_panes = 'B5'
    return esperado


def paragrafos(ws, linha, textos, largura=120, **kw):
    for t in textos:
        ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=3)
        c = ws.cell(linha, 2, t)
        c.font = fonte(**kw)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[linha].height = 14 * math.ceil(len(t) / largura) + 3
        linha += 1
    return linha


def reais(v):
    return f'{v:,.0f}'.replace(',', '.')


def aba_leia_me(wb, lista, meses, D):
    ws = wb.active
    ws.title = 'Leia-me'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['C'].width = 96
    p, u = rotulo(*meses[0]), rotulo(*meses[-1])
    antes, depois = rotulo(*anterior(*meses[0])), rotulo(*seguinte(*meses[-1]))
    ws['B2'] = 'Trattoria Bela Serra: lançamentos financeiros'
    ws['B2'].font = fonte(bold=True, size=16, color=TEAL)
    ws['B3'] = f'Empresa fictícia · DRE e fluxo de caixa de {p} a {u} · {reais(len(lista))} lançamentos'
    ws['B3'].font = fonte(color=SUAVE)
    linha = paragrafos(ws, 5, ['Material de uso exclusivo do processo seletivo da Zelo Gestão Financeira. A empresa, as pessoas '
                               'e os números são fictícios. Não compartilhe este arquivo.'], bold=True, color='9A6700')
    linha += 1
    ws.cell(linha, 2, 'Como os lançamentos formam os números da página').font = fonte(bold=True, size=12, color=TEAL)
    linha = paragrafos(ws, linha + 1, [
        f'• Resultado (DRE): some o Valor por Linha do DRE e Mês de competência, de {p} a {u}. '
        'A aba DRE faz isso com fórmulas.',
        f'• Fluxo de caixa: some o Valor por Linha do fluxo de caixa e Mês de caixa, de {p} a {u}. A aba Fluxo de caixa '
        f'faz isso com fórmulas. O saldo no banco em {dt.date(*meses[0], 1):%d/%m/%y} era de R$ {reais(D["caixa"][0]["inicial"])}.',
        '• Valor positivo é receita ou entrada de dinheiro; negativo é despesa ou saída.',
        f'• Há lançamentos de {antes}: vendas no crédito, compras, Simples e folha daquele mês que só foram pagos em {p}. '
        f'Eles entram no fluxo de caixa de {p}, mas não no DRE.',
        f'• Os lançamentos de {u} que vencem em {depois} estão em aberto (A receber ou A pagar). Entram no DRE de {u}, '
        'mas ainda não passaram pelo banco.',
        '• As taxas da maquininha e do iFood são descontadas no repasse. Por isso, no fluxo de caixa, entram em '
        'Recebimentos do salão e em Repasses do iFood, como na página.',
        '• A depreciação não mexe no caixa. Empréstimo, amortização, reforma e retiradas dos sócios mexem no caixa, '
        'mas não são receita nem despesa. Da parcela do empréstimo, só os juros entram no DRE.',
    ])
    linha += 1
    ws.cell(linha, 2, 'Colunas da aba Lançamentos').font = fonte(bold=True, size=12, color=TEAL)
    linha += 1
    for titulo, _, _, _ in COLUNAS:
        a = ws.cell(linha, 2, titulo)
        b = ws.cell(linha, 3, EXPLICA_COLUNAS[titulo])
        a.font, b.font = fonte(bold=True), fonte()
        a.alignment = Alignment(vertical='top')
        b.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[linha].height = 14 * math.ceil(len(EXPLICA_COLUNAS[titulo]) / 95) + 3
        linha += 1


def grava_resultados(caminho, esperado_por_aba):
    """Grava junto de cada fórmula o resultado já calculado, para quem abre o arquivo sem recalcular
    (pré-visualização no celular e no e-mail). Excel, LibreOffice e Google Planilhas recalculam ao abrir."""
    with zipfile.ZipFile(caminho) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    for indice, esperado in esperado_por_aba.items():
        nome = f'xl/worksheets/sheet{indice}.xml'
        troca = lambda m: f'{m.group(1)}<v>{esperado[m.group(2)] / 100:.15g}</v></c>'
        xml, feitas = re.subn(r'(<c r="([A-Z]+\d+)"[^>]*><f>[^<]*</f>)<v\s*/></c>', troca, partes[nome].decode('utf-8'))
        if feitas != len(esperado):
            falha(f'gravei {feitas} de {len(esperado)} resultados de fórmula em {nome}')
        partes[nome] = xml.encode('utf-8')
    with zipfile.ZipFile(caminho, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, b in partes.items():
            z.writestr(n, b)


def grava_xlsx(lista, meses, D):
    wb = Workbook()
    wb.properties.creator = 'Zelo Gestão Financeira'
    wb.properties.title = 'Trattoria Bela Serra: lançamentos financeiros'
    aba_leia_me(wb, lista, meses, D)
    aba_lancamentos(wb, lista)
    esperado = {   # número da aba no arquivo: resultados das fórmulas
        3: aba_resumo(wb, 'DRE', 'Resultado (DRE)',
                      'Soma do Valor da aba Lançamentos por Linha do DRE e Mês de competência.',
                      DRE_ABA, meses, lista, 'mes_comp', 'linha_dre'),
        4: aba_resumo(wb, 'Fluxo de caixa', 'Fluxo de caixa',
                      'Soma do Valor da aba Lançamentos por Linha do fluxo de caixa e Mês de caixa.',
                      FLUXO_ABA, meses, lista, 'mes_caixa', 'linha_fluxo', saldo_inicial=D['caixa'][0]['inicial']),
    }
    wb.calculation.fullCalcOnLoad = True
    ARQ_XLSX.parent.mkdir(exist_ok=True)
    wb.save(ARQ_XLSX)
    grava_resultados(ARQ_XLSX, esperado)


def grava_csv(lista):
    def texto(campo, v):
        if v is None:
            return ''
        if campo == 'valor':
            return f'{v / 100:.2f}'.replace('.', ',')
        if campo in ('mes_comp', 'mes_caixa'):
            return v.strftime('%Y-%m')
        if isinstance(v, dt.date):
            return v.strftime('%d/%m/%Y')
        return str(v)
    with open(ARQ_CSV, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f, delimiter=';', lineterminator='\r\n')
        w.writerow([t for t, *_ in COLUNAS])
        for l in lista:
            w.writerow([texto(campo, l[campo]) for _, campo, _, _ in COLUNAS])


def main():
    D = le_pagina()
    meses, lista = monta(D)
    confere(D, meses, lista)
    grava_xlsx(lista, meses, D)
    grava_csv(lista)
    abertos = [l for l in lista if l['situacao'] in ('A receber', 'A pagar')]
    print(f'{len(lista)} lançamentos conferidos com o DRE e o fluxo de caixa da página '
          f'({len(abertos)} em aberto em {FIM:%d/%m/%Y}).')
    print(f'Gravados: {ARQ_XLSX.relative_to(RAIZ)} e {ARQ_CSV.relative_to(RAIZ)}')


if __name__ == '__main__':
    main()
